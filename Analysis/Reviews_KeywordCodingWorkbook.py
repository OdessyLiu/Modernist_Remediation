"""Combine the first 100 ranked rows of each keyword/term CSV for human coding.

Requires openpyxl. Source CSVs are read only; existing workbooks are not replaced.
"""
import argparse
import csv
from itertools import islice
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "Results/Keywords_Rating"
N_VALUES = (1, 10, 100, 1000)
CODING_COLUMNS = ["selection", "codes", "notes"]
INTEGER_COLUMNS = {"line_number", "focus_count", "reference_count", "focus_docf", "reference_docf"}
TEXT_COLUMNS = {"item", "cql_query", *CODING_COLUMNS}
BLUE = "203E56"
TEAL = "DDEFEA"


def source_rows(path):
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        rows = list(islice(reader, 100))
        headers = reader.fieldnames
    if len(rows) != 100 or [int(row["line_number"]) for row in rows] != list(range(1, 101)):
        raise ValueError(f"Expected ranks 1–100 in {path.name}")
    scores = [float(row["keyness_score"]) for row in rows]
    if scores != sorted(scores, reverse=True):
        raise ValueError(f"Source is not ranked by descending keyness: {path.name}")
    return headers, rows


def text_cell(cell, value):
    cell.value = value
    cell.data_type = "s"


def build_workbook(input_dir, output):
    if output.exists():
        raise FileExistsError(f"Workbook already exists; choose another --output to preserve human coding: {output}")
    wb = Workbook()
    guide = wb.active
    guide.title = "Guide"
    guide.sheet_view.showGridLines = False
    guide.merge_cells("A1:E1")
    guide["A1"] = "Rating keywords and key terms | Top 100 for coding"
    guide["A1"].font = Font(size=18, bold=True, color=BLUE)
    guide.row_dimensions[1].height = 32
    instructions = [
        "16 result sheets, each containing the first 100 rows in its source CSV's keyness ranking (1,600 entries total).",
        "High sheets compare GB_HighRating against GB_LowRating; Low sheets reverse that comparison.",
        "N = 1, 10, 100, 1000 is the Simple Maths smoothing parameter. Repeated items across sheets remain separate.",
        "Enter your decisions in selection, codes, and notes (pale green columns). All coding cells start blank; no categories are imposed.",
        "All 16 original CSV columns are retained. Filter with header arrows; rank and item stay visible while scrolling.",
        "Relative frequencies are per million tokens; relative docf values are percentages, stored on a 0–100 scale.",
        "Save your coded copy under a new filename. The build script refuses to overwrite an existing workbook.",
        "See README.md and extraction_settings.json beside this workbook for corpus details and extraction settings.",
    ]
    for row_number, instruction in enumerate(instructions, 3):
        guide.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=5)
        guide.cell(row_number, 1, instruction).alignment = Alignment(wrap_text=True, vertical="center")
        guide.row_dimensions[row_number].height = 32
    guide.append([])
    index_header = 12
    for column, label in enumerate(("Open result sheet", "Focus", "Reference", "N", "Rows"), 1):
        cell = guide.cell(index_header, column, label)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color="FFFFFF", bold=True)
    for column, width in {"A": 38, "B": 25, "C": 25, "D": 12, "E": 12}.items():
        guide.column_dimensions[column].width = width
    guide.freeze_panes = "A13"
    expected = {}
    for label, reference in (("High", "Low"), ("Low", "High")):
        for kind in ("keywords", "keyterms"):
            for n in N_VALUES:
                path = input_dir / f"{label}_{kind}_N={n}.csv"
                headers, rows = source_rows(path)
                if headers[:2] != ["line_number", "item"]:
                    raise ValueError(f"Unexpected column layout: {path.name}")
                columns = headers[:2] + CODING_COLUMNS + headers[2:]
                ws = wb.create_sheet(path.stem)
                ws.sheet_view.showGridLines = False
                ws.sheet_view.zoomScale = 85
                ws.sheet_properties.tabColor = "367C80" if label == "High" else "AD7243"
                for row_number, value in enumerate((
                    f"{label} rating · {kind} · N = {n}",
                    f"Focus: GB_{label}Rating | Reference: GB_{reference}Rating | Source: {path.name}",
                    "Top 100 by source rank. Enter selection, codes, and notes in the green columns.",
                ), 1):
                    ws.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=len(columns))
                    ws.cell(row_number, 1, value)
                    ws.row_dimensions[row_number].height = 25
                ws["A1"].font = Font(size=17, bold=True, color=BLUE)
                ws["A3"].hyperlink = "#'Guide'!A1"
                ws["A3"].font = Font(color="367C80", italic=True)
                for col, name in enumerate(columns, 1):
                    cell = ws.cell(5, col, name)
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="367C80" if name in CODING_COLUMNS else BLUE)
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                    width = 21
                    if name == "line_number": width = 10
                    elif name == "item": width = 32
                    elif name == "notes": width = 48
                    elif name == "codes": width = 28
                    elif "relative" in name: width = 26
                    elif name == "cql_query": width = 50
                    ws.column_dimensions[cell.column_letter].width = width
                ws.row_dimensions[5].height = 62
                for row_number, row in enumerate(rows, 6):
                    for col, name in enumerate(columns, 1):
                        cell = ws.cell(row_number, col)
                        if name in CODING_COLUMNS:
                            cell.fill = PatternFill("solid", fgColor=TEAL)
                        elif name in TEXT_COLUMNS:
                            text_cell(cell, row[name])
                        else:
                            cell.value = int(row[name]) if name in INTEGER_COLUMNS else float(row[name])
                            cell.number_format = "#,##0" if name in INTEGER_COLUMNS else "0.000" if name == "keyness_score" else "0.00000"
                        cell.alignment = Alignment(vertical="top", wrap_text=name in TEXT_COLUMNS)
                    ws.row_dimensions[row_number].height = 30
                table = Table(displayName=f"Results_{len(expected) + 1:02d}", ref=f"A5:S105")
                table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
                ws.add_table(table)
                ws.freeze_panes = "C6"
                ws.print_title_rows = "1:5"
                ws.sheet_properties.pageSetUpPr.fitToPage = True
                ws.page_setup.orientation = "landscape"
                ws.page_setup.fitToWidth = 1
                ws.page_setup.fitToHeight = 0
                index_row = index_header + len(expected) + 1
                for col, value in enumerate((ws.title, f"GB_{label}Rating", f"GB_{reference}Rating", n, len(rows)), 1):
                    guide.cell(index_row, col, value)
                guide.cell(index_row, 1).hyperlink = f"#'{ws.title}'!A1"
                guide.cell(index_row, 1).style = "Hyperlink"
                expected[ws.title] = (columns, rows)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    # Verify the actual saved workbook, including blank coding cells and source values.
    import math
    saved = load_workbook(output)
    if saved.sheetnames != ["Guide", *expected]:
        raise ValueError("Saved sheet list differs from the expected 16 result sets")
    for name, (columns, rows) in expected.items():
        ws = saved[name]
        assert ws.max_row == 105 and ws.max_column == len(columns) == 19
        assert [cell.value for cell in ws[5]] == columns
        for row_number, source in enumerate(rows, 6):
            for col, field in enumerate(columns, 1):
                value = ws.cell(row_number, col).value
                if field in CODING_COLUMNS:
                    assert value is None
                elif field in TEXT_COLUMNS:
                    assert value == source[field] and ws.cell(row_number, col).data_type == "s"
                else:
                    assert math.isclose(value, float(source[field]), rel_tol=1e-14, abs_tol=1e-14)
        assert ws.freeze_panes == "C6" and len(ws.tables) == 1
    saved.close()
    print(f"Saved and verified: {output}\n16 result sheets × 100 rows = 1,600 entries, plus Guide.")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, default=RESULTS)
    parser.add_argument("--output", type=Path, default=RESULTS / "Keywords_Rating_Top100_Coding.xlsx")
    args = parser.parse_args()
    build_workbook(args.input_dir, args.output)


if __name__ == "__main__":
    main()
